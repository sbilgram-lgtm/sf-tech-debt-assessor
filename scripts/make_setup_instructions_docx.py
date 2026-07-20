"""
Generates SF_Tech_Debt_Assessor_Setup_Instructions_2026-06-22.docx
Customer-facing setup guide: Connected Apps, External Client Apps, Render, Docker.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "/Users/sbilgram/Desktop/SF_Tech_Debt_Assessor_Setup_Instructions_2026-06-22.docx"

BLUE   = RGBColor(0x03, 0x2D, 0x60)
DARK   = RGBColor(0x2C, 0x3E, 0x50)
GREY   = RGBColor(0x7F, 0x8C, 0x8D)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x27, 0xAE, 0x60)
RED    = RGBColor(0xC0, 0x39, 0x2B)
AMBER  = RGBColor(0xD3, 0x54, 0x00)


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_heading(doc, text, level=1, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.bold = True
    if level == 1:
        r.font.size = Pt(20)
        r.font.color.rgb = color or BLUE
    elif level == 2:
        r.font.size = Pt(14)
        r.font.color.rgb = color or BLUE
        p.paragraph_format.space_before = Pt(14)
    else:
        r.font.size = Pt(11)
        r.font.color.rgb = color or DARK
        p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = color or DARK
    p.paragraph_format.space_after = Pt(4)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(f'ℹ  {text}')
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = GREY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_warning(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(f'⚠  {text}')
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = AMBER
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_step(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(3)
    run_num = p.add_run(f'{number}.  ')
    run_num.font.bold = True
    run_num.font.size = Pt(10.5)
    run_num.font.color.rgb = BLUE
    run_text = p.add_run(text)
    run_text.font.size = Pt(10.5)
    run_text.font.color.rgb = DARK


def add_code(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x18, 0x56, 0x37)
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_table(doc, headers, rows, header_bg='032D60'):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = 'Table Grid'
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = WHITE
        set_cell_bg(hdr_cells[i], header_bg)
    for idx, row_data in enumerate(rows):
        row = tbl.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = val
            row[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        if idx % 2 == 1:
            for cell in row:
                set_cell_bg(cell, 'F4F6F7')
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


# ── Build document ────────────────────────────────────────────────────

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.2)
    section.right_margin  = Cm(2.2)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)

# ── Cover ─────────────────────────────────────────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('SF Tech Debt Assessor')
r.font.size = Pt(28)
r.font.bold = True
r.font.color.rgb = BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Setup & User Guide  —  June 22, 2026')
r.font.size = Pt(13)
r.font.italic = True
r.font.color.rgb = GREY

by = doc.add_paragraph()
by.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = by.add_run('By Steven Bilgram, Success Architect  |  sf-tech-debt-assessor.onrender.com')
r.font.size = Pt(10)
r.font.color.rgb = GREY

doc.add_paragraph()

intro = doc.add_paragraph()
r = intro.add_run(
    'This guide explains how to register the SF Tech Debt Assessor in your Salesforce org and run a full '
    'assessment. The app connects via OAuth, performs a read-only scan across 343 checks in 23 categories, '
    'and produces a scored report you can export as PDF, Excel, or CSV.'
)
r.font.size = Pt(10.5)
r.font.color.rgb = DARK
intro.paragraph_format.space_after = Pt(10)

doc.add_page_break()

# ── Section 1: Which setup do I need? ────────────────────────────────
add_heading(doc, '1. Which Setup Type Do I Need?', level=1)
add_body(doc,
    'The registration steps differ slightly depending on your Salesforce org type. '
    'Use the table below to identify your org type, then follow the matching instructions.'
)

add_table(doc,
    ['Org Type', 'Setup Type to Use', 'Works with Render', 'Works with Docker'],
    [
        ('Production / Sandbox / Developer Edition', 'Connected App  (Option A)', '✓', '✓'),
        ('Trailhead Playground / Spring \'25+ org', 'External Client App  (Option B)', '✓', '✗'),
    ]
)

add_note(doc,
    'How to tell: Go to Setup and search App Manager. '
    'If you see a "New Connected App" button, use Option A. '
    'If you only see External Client Apps, use Option B — but note that local Docker will not be available.'
)

doc.add_paragraph()

# ── Section 2: Option A — Connected App ──────────────────────────────
add_heading(doc, '2. Option A — Connected App (Production, Sandbox, Developer Edition)', level=1)
add_body(doc, 'Supports both the hosted Render version and local Docker.')

add_heading(doc, 'Step 1 — Create the Connected App', level=2)
add_step(doc, 1, 'Log in as an Administrator and go to Setup → App Manager → New Connected App.')
add_step(doc, 2, 'Fill in the required fields:')
add_code(doc, '    Connected App Name:  SF Tech Debt Assessor')
add_code(doc, '    API Name:            SF_Tech_Debt_Assessor  (auto-fills)')
add_code(doc, '    Contact Email:       your email address')
add_step(doc, 3, 'Check Enable OAuth Settings.')
add_step(doc, 4, 'In the Callback URL field, enter the URLs for the options you want — one per line:')
add_code(doc, '    Hosted Render:  https://sf-tech-debt-assessor.onrender.com/auth/callback')
add_code(doc, '    Local Docker:   http://localhost:3001/auth/callback')
add_note(doc, 'You can include both — Salesforce will match whichever URL is used at login time.')
add_step(doc, 5, 'Under Selected OAuth Scopes, add both of the following:')
add_code(doc, '    Access and manage your data (api)')
add_code(doc, '    Perform requests on your behalf at any time (refresh_token, offline_access)')
add_step(doc, 6, 'Uncheck "Require Proof Key for Code Exchange (PKCE)" — this must be disabled.')
add_step(doc, 7, 'Click Save → then Continue.')
add_step(doc, 8, 'Wait 10 minutes. Salesforce needs time to activate the app before it can be used.')

add_heading(doc, 'Step 2 — Copy your credentials', level=2)
add_step(doc, 1, 'Go back to App Manager, find your app, click the dropdown arrow → View.')
add_step(doc, 2, 'Click Manage Consumer Details (you may be asked to verify your identity).')
add_step(doc, 3, 'Copy the Consumer Key — this is your Client ID.')
add_step(doc, 4, 'Copy the Consumer Secret — this is your Client Secret.')

doc.add_paragraph()

# ── Section 3: Option B — External Client App ────────────────────────
add_heading(doc, '3. Option B — External Client App (Trailhead Playground / Spring \'25+ Orgs)', level=1)
add_warning(doc,
    'External Client Apps only support the hosted Render version. '
    'Local Docker requires HTTPS and is not supported with an http://localhost callback URL.'
)

add_heading(doc, 'Step 1 — Create the External Client App', level=2)
add_step(doc, 1, 'Log in as an Administrator and go to Setup → External Client Apps → New.')
add_step(doc, 2, 'Fill in the required fields:')
add_code(doc, '    Label:          SF Tech Debt Assessor')
add_code(doc, '    Name:           SF_Tech_Debt_Assessor  (auto-fills)')
add_code(doc, '    Contact Email:  your email address')
add_step(doc, 3, 'Under OAuth Settings, check Enable OAuth.')
add_step(doc, 4, 'Set the Callback URL to:')
add_code(doc, '    https://sf-tech-debt-assessor.onrender.com/auth/callback')
add_step(doc, 5, 'Under OAuth Scopes, add both of the following:')
add_code(doc, '    Access and manage your data (api)')
add_code(doc, '    Perform requests on your behalf at any time (refresh_token, offline_access)')
add_step(doc, 6, 'Uncheck "Require Proof Key for Code Exchange (PKCE)" if it appears — leave it disabled.')
add_step(doc, 7, 'Click Save. Wait 10 minutes for Salesforce to activate the app.')
add_step(doc, 8, 'Go back to the External Client App → click the app name → copy your Client ID and Client Secret.')

doc.add_paragraph()

# ── Section 4: Running the assessment (Render) ───────────────────────
add_heading(doc, '4. Running an Assessment — Hosted Version (Render)', level=1)
add_note(doc, 'The hosted site may take up to 30 seconds to load if it has been idle. This is normal.')

add_step(doc, 1, 'Open https://sf-tech-debt-assessor.onrender.com in your browser.')
add_step(doc, 2, 'Enter your credentials:')
add_code(doc, '    Org / Sandbox URL  — your org\'s My Domain URL (see examples below)')
add_code(doc, '    Client ID          — Consumer Key / Client ID from your app setup')
add_code(doc, '    Client Secret      — Consumer Secret / Client Secret from your app setup')
add_step(doc, 3, 'Click Connect to Salesforce. Log in to Salesforce when prompted, then click Allow.')
add_step(doc, 4, 'The assessment starts automatically. A progress indicator shows each category as it scans.')
add_step(doc, 5, 'Click any category panel to expand and see the individual findings.')
add_step(doc, 6, 'Click Show affected records on any finding to see the specific records causing the issue.')
add_step(doc, 7, 'Export your results using the buttons at the top right:')
add_code(doc, '    Export PDF     — full report with scores, findings, and affected records')
add_code(doc, '    Export Excel   — one tab per category plus a Summary tab')
add_code(doc, '    Export CSV     — flat file, one row per affected record')
add_code(doc, '    Remediation Roadmap — phased action plan (Critical → High → Medium → Low)')

doc.add_paragraph()
add_heading(doc, 'Org URL Format Examples', level=3)
add_table(doc,
    ['Org Type', 'Example URL'],
    [
        ('Production', 'https://yourcompany.my.salesforce.com'),
        ('Sandbox', 'https://yourcompany--sandboxname.sandbox.my.salesforce.com'),
        ('Developer Edition', 'https://yourname-dev-ed.develop.my.salesforce.com'),
        ('Trailhead Playground', 'https://yourname-dev-ed.trailblaze.my.salesforce.com'),
    ]
)
add_note(doc,
    'Find your exact org URL in Salesforce → Setup → Company Settings → My Domain. '
    'Use everything up to and including .salesforce.com with no trailing slash.'
)

doc.add_page_break()

# ── Section 5: Docker ─────────────────────────────────────────────────
add_heading(doc, '5. Running with Docker (Local / On-Premise)', level=1)
add_body(doc,
    'Docker lets you run the app inside your own network without depending on Render. '
    'Use this option when you need to keep data entirely on-premises.'
)
add_warning(doc,
    'Docker requires a Connected App (Option A), not an External Client App. '
    'External Client Apps only allow HTTPS callback URLs and do not support http://localhost.'
)

add_heading(doc, 'Prerequisites', level=2)
add_body(doc, '• Docker Desktop installed — download from https://www.docker.com/products/docker-desktop/')
add_body(doc, '• A Connected App with http://localhost:3001/auth/callback added to its Callback URL list')

add_heading(doc, 'Windows / Linux / Intel Mac', level=3)
add_code(doc, 'docker pull ghcr.io/sbilgram-lgtm/sf-tech-debt-assessor:latest')
doc.add_paragraph()
add_code(doc, 'docker run -d \\')
add_code(doc, '  -p 3001:3001 \\')
add_code(doc, '  -e SESSION_SECRET=change-this-to-a-random-string \\')
add_code(doc, '  -e NODE_ENV=production \\')
add_code(doc, '  --name sf-assessor \\')
add_code(doc, '  ghcr.io/sbilgram-lgtm/sf-tech-debt-assessor:latest')

add_heading(doc, 'Apple Silicon Mac (M1 / M2 / M3) — Build Locally', level=3)
add_note(doc, 'The published image is Intel (amd64) only. Apple Silicon Macs must build the image locally.')
add_code(doc, 'git clone https://github.com/sbilgram-lgtm/sf-tech-debt-assessor')
add_code(doc, 'cd sf-tech-debt-assessor')
add_code(doc, 'docker build -t sf-tech-debt-assessor .')
doc.add_paragraph()
add_code(doc, 'docker run -d \\')
add_code(doc, '  -p 3001:3001 \\')
add_code(doc, '  -e SESSION_SECRET=change-this-to-a-random-string \\')
add_code(doc, '  -e NODE_ENV=production \\')
add_code(doc, '  --name sf-assessor \\')
add_code(doc, '  sf-tech-debt-assessor')

add_body(doc, 'Then open http://localhost:3001 in your browser.')

doc.add_page_break()

# ── Section 6: Troubleshooting ────────────────────────────────────────
add_heading(doc, '6. Troubleshooting Common Errors', level=1)

add_table(doc,
    ['Error Message', 'Cause', 'Fix'],
    [
        ('Authentication failed', 'Wrong Client ID or Client Secret',
         'Copy-paste directly from Manage Consumer Details — no extra spaces'),
        ('redirect_uri_mismatch', 'Callback URL doesn\'t exactly match what\'s in your app',
         'Check the URL is entered exactly as shown in Section 2 or 3, then Save and wait 10 minutes'),
        ('invalid_client_id', 'App not yet active',
         'Wait the full 10 minutes after creating the app, then try again'),
        ('missing required code challenge', 'PKCE is still enabled',
         'Go back to your app in Setup and uncheck Require PKCE'),
        ('no matching manifest for linux/arm64  (Docker)',
         'GHCR image is Intel-only',
         'Follow the Apple Silicon build-locally instructions in Section 5'),
        ('port is already allocated  (Docker)', 'Port 3001 is in use',
         'Stop the existing container or change the port mapping to -p 3002:3001'),
    ],
    header_bg='1A5276'
)

doc.add_paragraph()

# ── Section 7: Permissions ────────────────────────────────────────────
add_heading(doc, '7. Required Salesforce Permissions', level=1)
add_body(doc,
    'The user who authenticates must have the following permissions. '
    'System Administrator profile grants all of them automatically.'
)
add_body(doc, '• API Enabled (on their Profile)')
add_body(doc, '• View Setup and Configuration')
add_body(doc, '• Modify Metadata Through Metadata API Functions (for full results)')

doc.add_paragraph()

# ── Section 8: What the app checks ───────────────────────────────────
add_heading(doc, '8. What the App Checks — 343 Checks in 23 Categories', level=1)
add_body(doc,
    'The assessment runs a read-only scan and scores each category from 0 to 100. '
    'Nothing is written back to Salesforce. Results exist only in your browser session.'
)

add_table(doc,
    ['Category', 'Checks', 'Group'],
    [
        ('Configuration', '13', 'Configuration & Architecture'),
        ('Code Quality', '46', 'Code & Development'),
        ('Data Model', '4', 'Configuration & Architecture'),
        ('Service Cloud', '69', 'CRM & Service'),
        ('Sharing & Security', '24', 'Security & Access'),
        ('Integrations', '9', 'Configuration & Architecture'),
        ('Test Coverage', '7', 'Code & Development'),
        ('Org Limits', '5', 'Performance & Limits'),
        ('Duplicate & Matching Rules', '4', 'Configuration & Architecture'),
        ('Reports & Dashboards', '3', 'Governance & Hygiene'),
        ('Email Templates', '3', 'Governance & Hygiene'),
        ('Platform Events & CDC', '3', 'Performance & Limits'),
        ('Managed Packages', '3', 'CRM & Service'),
        ('Custom Metadata & Settings', '3', 'Configuration & Architecture'),
        ('Record Types & Page Layouts', '4', 'Configuration & Architecture'),
        ('Einstein & AI', '9', 'CRM & Service'),
        ('Experience Cloud', '15', 'Security & Access'),
        ('Connected App Security', '12', 'Security & Access'),
        ('LWC & Components', '39', 'Code & Development'),
        ('OmniStudio', '26', 'Code & Development'),
        ('Performance', '22', 'Performance & Limits'),
        ('Notes & Attachments', '12', 'Governance & Hygiene'),
        ('Flow Quality', '8', 'Code & Development'),
    ]
)

# ── Save ──────────────────────────────────────────────────────────────
doc.save(OUTPUT)
print(f'Saved: {OUTPUT}')
